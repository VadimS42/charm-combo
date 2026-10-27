"""Charm Combo: build the voice actor's script from data/characters.json.

Usage (from the project folder):
    python3 tools/voice_script/make_voice_script.py [out_dir]
Writes <out_dir>/Charm_Combo_Voice_Script.xlsx (the recording sheet: one tab
per character, one row per line) and <out_dir>/voice_lines.json (the same
lines, which make_voice_docx.js turns into the readable .docx script).
Default out_dir: docs/voice. Needs Python 3 with openpyxl.

Every spoken line in characters.json carries an "id" (the file name to record
it as) and a "d" delivery note; see docs/EDITING_DIALOGUE.md.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "docs", "voice")
SHEET = {"succubus": "Amara", "mermaid": "Marina", "slime": "Opal", "dryad": "Hazel"}


def collect(ch):
    """All voiced lines of one character, in play order, with scene + context."""
    rows = []

    def block(lines, scene, context=""):
        prev = context
        for ln in lines:
            if ln.get("k") == "slot":
                continue  # a writer's note: where the user's own lines go
            if ln.get("k") == "n":
                prev = ln["t"]
                continue
            rows.append({"id": ln["id"], "scene": scene, "context": prev, "line": ln["t"],
                         "emotion": ln.get("e", "neutral"), "note": ln.get("d", "")})
            prev = "Her previous line: " + ln["t"]

    s = ch["story"]
    block(s["intro"], "Meeting her")
    for i, o in enumerate(s["introChoice"]["options"], 1):
        block(o["resp"], "Meeting her: reply %d" % i, "You said: " + o["label"])
    for n, lv in enumerate(s["levels"], 1):
        last = n == len(s["levels"])
        block(lv["win"], "After chapter %d%s" % (n, " (last scene)" if last else ""), "Chapter %d's puzzle has just been won." % n)
        if lv.get("choice"):
            for i, o in enumerate(lv["choice"]["options"], 1):
                block(o["resp"], "After chapter %d: reply %d" % (n, i), "You said: " + o["label"])
    night = s.get("night")
    if isinstance(night, dict):
        block(night.get("lines", []), "Her night scene (full-screen CG)", "Her last chapter's scene has just ended. The two of you are alone.")
    for key, label in (("warmth", "Warm ending"), ("spark", "Teasing ending")):
        ep = s["endings"][key]["epilogue"]
        block(ep["intro"] + ep["middle"] + ep["choice_lines"], label + ": the morning after", "The night after her last chapter.")
        for i, o in enumerate(ep["choice"]["options"], 1):
            block(o["resp"], label + ": final reply %d" % i, "You said: " + o["label"])
    block(ch["giftLine"], "Gift scene (%s)" % ch["gift"]["name"])
    for field, scene in (("favQuips", "Puzzle quip: you matched her favourite gem"),
                         ("leastQuips", "Puzzle quip: you matched a gem she dislikes"),
                         ("chatQuips", "Puzzle quip: chatter while you play")):
        for q in ch[field]:
            if isinstance(q, dict):
                rows.append({"id": q["id"], "scene": scene, "context": "Short bark, shown as a pop-up while you play.",
                             "line": q["t"], "emotion": q.get("e", "neutral"), "note": q.get("d", "")})
    return rows


def main():
    os.makedirs(OUT, exist_ok=True)
    chars = json.load(open(os.path.join(ROOT, "data", "characters.json"), encoding="utf-8"))
    data = []
    for ch in chars:
        data.append({"id": ch["id"], "name": ch["name"], "sheet": SHEET.get(ch["id"], ch["name"].split()[0]),
                     "role": ch.get("role", ""), "voice": ch.get("voice", ""), "accent": ch.get("c", "#888888"),
                     "pronunciation": ch.get("pronunciation", []), "lines": collect(ch), "story": ch["story"],
                     "giftLine": ch["giftLine"], "gift": ch["gift"],
                     "quips": {f: ch[f] for f in ("favQuips", "leastQuips", "chatQuips")}})
    json.dump(data, open(os.path.join(OUT, "voice_lines.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)

    wb = Workbook()
    F = "Arial"
    thin = Side(style="thin", color="D9C9D3")
    wrap = Alignment(wrap_text=True, vertical="top")

    # --- Read me
    ws = wb.active
    ws.title = "Read me"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    r = 1
    ws.cell(r, 1, "Charm Combo · Voice script").font = Font(name=F, size=18, bold=True, color="8E2F63")
    r += 1
    ws.cell(r, 1, "An 18+ match-3 romance. One tab per character; one row per line to record.").font = Font(name=F, size=11, color="6B5466")
    r += 2
    info = [
        ("What to record", "Every row on the character tabs. The Line column is the exact text; Emotion and Delivery note say how; What just happened gives the moment. Narration is never recorded."),
        ("File names", "Save each line as its File name (the Line ID plus .wav), one line per file. Keep the IDs exactly as written: they match the lines in the game's files, so the recordings can be hooked up later without renaming."),
        ("Audio", "WAV, 48 kHz, 24-bit, mono. About half a second of silence before and after each line. No music or effects."),
        ("Takes", "Deliver your best take per line. If you record alternates, add _alt1, _alt2 to the file name and mention it in Notes."),
        ("Content", "Romance with sexual content. All characters are adults. Each girl has a night scene with the player; its lines in this script are intimate but not graphic. More lines for those scenes may come later as a separate session."),
        ("Your columns", "Fill in Done (Yes) and Notes (pick-ups, alternates, questions) on each character tab. Example: Done = Yes, Notes = \"alt1 is softer\"."),
        ("Emotion", "The portrait the game shows while the line is on screen: neutral, happy, shy, surprised, teasing, sad or angry."),
    ]
    top = Alignment(vertical="top")
    for k, v in info:
        ws.cell(r, 1, k).font = Font(name=F, size=11, bold=True)
        ws.cell(r, 1).alignment = top
        c = ws.cell(r, 2, v)
        c.font = Font(name=F, size=11)
        c.alignment = wrap
        ws.row_dimensions[r].height = 15 * (len(v) // 95 + 1)
        r += 1
    r += 1
    ws.cell(r, 1, "Lines per character").font = Font(name=F, size=13, bold=True, color="8E2F63")
    r += 1
    first_count = r
    for d in data:
        ws.cell(r, 1, d["name"]).font = Font(name=F, size=11)
        c = ws.cell(r, 2, "=COUNTA('%s'!B6:B500)" % d["sheet"])
        c.font = Font(name=F, size=11)
        c.alignment = Alignment(horizontal="left")
        r += 1
    ws.cell(r, 1, "Total").font = Font(name=F, size=11, bold=True)
    c = ws.cell(r, 2, "=SUM(B%d:B%d)" % (first_count, r - 1))
    c.font = Font(name=F, size=11, bold=True)
    c.alignment = Alignment(horizontal="left")
    r += 2
    ws.cell(r, 1, "The characters").font = Font(name=F, size=13, bold=True, color="8E2F63")
    r += 1
    for d in data:
        ws.cell(r, 1, d["name"]).font = Font(name=F, size=11, bold=True, color=d["accent"].lstrip("#"))
        ws.cell(r, 1).alignment = top
        c = ws.cell(r, 2, d["voice"])
        c.font = Font(name=F, size=11)
        c.alignment = wrap
        ws.row_dimensions[r].height = 15 * (len(d["voice"]) // 95 + 1)
        r += 1
        if d["pronunciation"]:
            ws.cell(r, 1, "  Say it").font = Font(name=F, size=10, italic=True, color="6B5466")
            ws.cell(r, 1).alignment = top
            text = " · ".join("%s = %s" % (a, b) for a, b in d["pronunciation"])
            c = ws.cell(r, 2, text)
            c.font = Font(name=F, size=10, italic=True, color="6B5466")
            c.alignment = wrap
            ws.row_dimensions[r].height = 14 * (len(text) // 100 + 1)
            r += 1
    ws.cell(r + 1, 1, "Made from data/characters.json by tools/voice_script/make_voice_script.py.").font = Font(name=F, size=9, color="8A7A86")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # --- one tab per character
    cols = [("#", 5), ("Line ID", 20), ("File name", 24), ("Scene", 24), ("What just happened", 42),
            ("Line", 60), ("Emotion", 11), ("Delivery note", 40), ("Done", 8), ("Notes", 28)]
    for d in data:
        ws = wb.create_sheet(d["sheet"])
        ws.sheet_view.showGridLines = False
        acc = d["accent"].lstrip("#")
        ws.cell(1, 1, "%s · %s" % (d["name"], d["role"])).font = Font(name=F, size=16, bold=True, color=acc)
        c = ws.cell(2, 1, d["voice"])
        c.font = Font(name=F, size=10, color="4A3A47")
        c.alignment = wrap
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
        ws.row_dimensions[2].height = 44
        if d["pronunciation"]:
            ws.cell(3, 1, "Say it: " + " · ".join("%s = %s" % (a, b) for a, b in d["pronunciation"])).font = Font(name=F, size=10, italic=True, color="6B5466")
        ws.cell(4, 1, "Lines:").font = Font(name=F, size=10, bold=True)
        ws.cell(4, 2, "=COUNTA(B6:B500)").font = Font(name=F, size=10, bold=True)
        ws.cell(4, 2).alignment = Alignment(horizontal="left")
        ws.cell(4, 3, "Done:").font = Font(name=F, size=10, bold=True)
        ws.cell(4, 4, '=COUNTIF(I6:I500,"Yes")').font = Font(name=F, size=10, bold=True)
        ws.cell(4, 4).alignment = Alignment(horizontal="left")
        hdr = 5
        for i, (name, w) in enumerate(cols, 1):
            c = ws.cell(hdr, i, name)
            c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=acc)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = Border(bottom=thin)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(hdr + 1, 3)
        band = PatternFill("solid", fgColor="FBF4F8")
        last_scene = None
        for n, ln in enumerate(d["lines"], 1):
            rr = hdr + n
            vals = [n, ln["id"], ln["id"] + ".wav", ln["scene"], ln["context"], ln["line"], ln["emotion"], ln["note"], None, None]
            for i, v in enumerate(vals, 1):
                c = ws.cell(rr, i, v)
                c.font = Font(name=F, size=10, bold=(i == 6), color=("8A7A86" if i in (1, 3, 5) else "2B2330"))
                c.alignment = wrap
                c.border = Border(bottom=thin)
                if ln["scene"] != last_scene and i == 4:
                    c.font = Font(name=F, size=10, bold=True, color="2B2330")
            if n % 2 == 0:
                for i in range(1, len(cols) + 1):
                    ws.cell(rr, i).fill = band
            last_scene = ln["scene"]
            longest = max(len(ln["line"]) / 55.0, len(ln["context"]) / 38.0, len(ln["note"]) / 36.0)
            ws.row_dimensions[rr].height = max(15, 13.5 * (int(longest) + 1))
        dv = DataValidation(type="list", formula1='"Yes"', allow_blank=True)
        dv.add("I%d:I%d" % (hdr + 1, hdr + len(d["lines"])))
        ws.add_data_validation(dv)
        ws.auto_filter.ref = "A%d:J%d" % (hdr, hdr + len(d["lines"]))
        ws.print_title_rows = "%d:%d" % (hdr, hdr)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    path = os.path.join(OUT, "Charm_Combo_Voice_Script.xlsx")
    wb.save(path)
    print("wrote", path, "and voice_lines.json;", ", ".join("%s %d" % (d["sheet"], len(d["lines"])) for d in data))


if __name__ == "__main__":
    main()
